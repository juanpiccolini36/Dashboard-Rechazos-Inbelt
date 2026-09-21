"""
agregar_dropdown_tipo_falla.py
------------------------------------------------------------
Agrega una lista desplegable a la columna "Tipo de falla" de cada
pestaña "Rechazos *", usando como fuente la MISMA lista de sectores
que se usa para generar las pestañas (columna M de la Matriz
General, ya normalizada).

La lista fuente se guarda en una hoja de referencia
"Config_Tipos_Falla" - si en el futuro se agrega un sector nuevo,
alcanza con agregarlo ahi una vez (actualizar_rechazos.py lo hace
solo cuando crea una pestaña de sector nueva) y el desplegable de
TODAS las pestañas ya lo va a ofrecer, sin tocar cada una a mano.

USO:
    python agregar_dropdown_tipo_falla.py --rechazos ARCHIVO.xlsx --salida ARCHIVO_NUEVO.xlsx
"""

import argparse
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font

FILA_LIMITE_VALIDACION = 5000
NOMBRE_HOJA_LISTA = "Config_Tipos_Falla"

# misma lista y mismo orden que generan las pestañas de sector
# (ver normalizar_nombre_sector en actualizar_rechazos.py)
SECTORES = [
    "Compras", "Gen. Prog. Produccion", "Granallado", "Herreria",
    "Ingenieria", "Inyeccion", "Laboratorio", "Mantenimiento",
    "Mec. Torno", "Mecanizado", "Mesa de Corte", "Moldeo", "Pañol",
    "Poliuretano", "RRHH", "Recepcion", "Revestimiento", "Servicio",
    "Terminacion",
]


def asegurar_hoja_lista(wb):
    if NOMBRE_HOJA_LISTA in wb.sheetnames:
        return wb[NOMBRE_HOJA_LISTA]
    ws = wb.create_sheet(NOMBRE_HOJA_LISTA)
    ws["A1"] = "Tipo de falla (lista fuente del desplegable - agregar sectores nuevos aca)"
    ws["A1"].font = Font(bold=True)
    ws.column_dimensions["A"].width = 55
    for i, sector in enumerate(SECTORES, start=2):
        ws.cell(row=i, column=1, value=sector)
    return ws


def aplicar_dropdown(ws):
    headers = [c.value for c in ws[1]]
    if "Tipo de falla" not in headers:
        return False
    col = headers.index("Tipo de falla") + 1
    letra = get_column_letter(col)

    dv = DataValidation(
        type="list",
        formula1=f"={NOMBRE_HOJA_LISTA}!$A$2:$A$200",
        allow_blank=True,
    )
    dv.error = "Elegi un tipo de falla de la lista"
    dv.errorTitle = "Valor invalido"
    ws.add_data_validation(dv)
    dv.add(f"{letra}2:{letra}{FILA_LIMITE_VALIDACION}")
    return True


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--rechazos", required=True)
    ap.add_argument("--salida", default=None)
    args = ap.parse_args()
    salida = args.salida or args.rechazos

    wb = load_workbook(args.rechazos)
    ws_lista = asegurar_hoja_lista(wb)
    print(f"Hoja '{NOMBRE_HOJA_LISTA}' lista, con {ws_lista.max_row - 1} sectores.")

    aplicadas = 0
    for nombre in wb.sheetnames:
        if not nombre.lower().startswith("rechazos"):
            continue
        ws = wb[nombre]
        if aplicar_dropdown(ws):
            aplicadas += 1
            print(f"  -> desplegable aplicado en '{nombre}'")

    wb.save(salida)
    print(f"\nTotal hojas con desplegable: {aplicadas}")
    print(f"Guardado: {salida}")


if __name__ == "__main__":
    main()
