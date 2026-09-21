"""
aplicar_formato.py
------------------------------------------------------------
Copia el formato visual (ancho de columnas, estilo de Tabla de
Excel, panel inmovilizado) de UNA pestaña de referencia ya
formateada a mano, hacia todas las demas pestañas "Rechazos *".

Empareja las columnas por NOMBRE DE ENCABEZADO, no por posicion,
para no desalinear anchos si alguna hoja tiene una columna de mas
o de menos.

USO:
    python aplicar_formato.py --rechazos ARCHIVO.xlsx --referencia "Rechazos Compras" --salida ARCHIVO_FORMATEADO.xlsx
"""

import argparse
import re
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


def leer_formato_referencia(ws_ref):
    headers = [c.value for c in ws_ref[1]]
    ancho_por_header = {}
    for i, h in enumerate(headers, start=1):
        letra = get_column_letter(i)
        dim = ws_ref.column_dimensions.get(letra)
        if dim and dim.width:
            ancho_por_header[h] = dim.width

    estilo = None
    if ws_ref.tables:
        tabla_ref = list(ws_ref.tables.values())[0]
        estilo = tabla_ref.tableStyleInfo
    freeze = ws_ref.freeze_panes
    return ancho_por_header, estilo, freeze


def nombre_tabla_valido(nombre_hoja):
    """Los nombres de Tabla de Excel no pueden tener espacios ni caracteres
    especiales, y deben ser unicos en el libro."""
    limpio = re.sub(r"[^A-Za-z0-9_]", "_", nombre_hoja)
    return f"Tabla_{limpio}"


def aplicar_formato_hoja(ws, ancho_por_header, estilo, freeze):
    headers = [c.value for c in ws[1]]

    # 1) anchos de columna, emparejados por nombre de encabezado
    for i, h in enumerate(headers, start=1):
        if h in ancho_por_header:
            letra = get_column_letter(i)
            ws.column_dimensions[letra].width = ancho_por_header[h]

    # 2) panel inmovilizado
    if freeze:
        ws.freeze_panes = freeze

    # 3) tabla de Excel con el mismo estilo
    if estilo is not None:
        # sacar cualquier tabla previa en esta hoja para no duplicar
        for nombre_existente in list(ws.tables.keys()):
            del ws.tables[nombre_existente]

        ultima_fila = ws.max_row
        ultima_col_letra = get_column_letter(len(headers))
        ref = f"A1:{ultima_col_letra}{ultima_fila}"
        tabla = Table(displayName=nombre_tabla_valido(ws.title), ref=ref)
        tabla.tableStyleInfo = TableStyleInfo(
            name=estilo.name,
            showFirstColumn=estilo.showFirstColumn,
            showLastColumn=estilo.showLastColumn,
            showRowStripes=estilo.showRowStripes,
            showColumnStripes=estilo.showColumnStripes,
        )
        ws.add_table(tabla)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--rechazos", required=True)
    ap.add_argument("--referencia", required=True, help="Nombre exacto de la pestaña ya formateada")
    ap.add_argument("--salida", default=None)
    args = ap.parse_args()
    salida = args.salida or args.rechazos

    wb = load_workbook(args.rechazos)
    if args.referencia not in wb.sheetnames:
        print(f"ERROR: no encuentro la pestaña '{args.referencia}' en el archivo.")
        return

    ws_ref = wb[args.referencia]
    ancho_por_header, estilo, freeze = leer_formato_referencia(ws_ref)
    print(f"Formato leido de '{args.referencia}': {len(ancho_por_header)} columnas con ancho, "
          f"estilo de tabla={'si' if estilo else 'no'}, freeze_panes={freeze}")

    aplicadas = 0
    for nombre in wb.sheetnames:
        if not nombre.lower().startswith("rechazos") or nombre == args.referencia:
            continue
        ws = wb[nombre]
        aplicar_formato_hoja(ws, ancho_por_header, estilo, freeze)
        aplicadas += 1
        print(f"  -> formato aplicado a '{nombre}'")

    wb.save(salida)
    print(f"\nTotal hojas formateadas: {aplicadas}")
    print(f"Guardado: {salida}")


if __name__ == "__main__":
    main()
