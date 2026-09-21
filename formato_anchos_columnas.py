"""
formato_anchos_columnas.py
------------------------------------------------------------
Fija el ancho de columnas puntuales en todas las pestañas
"Rechazos *", dejando el resto de las columnas sin tocar.

USO:
    python formato_anchos_columnas.py --rechazos ARCHIVO.xlsx --salida ARCHIVO_NUEVO.xlsx
"""

import argparse
from openpyxl import load_workbook

ANCHOS_FIJOS = {
    "D": 30,   # Proceso
    "F": 70,   # nombre
}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--rechazos", required=True)
    ap.add_argument("--salida", default=None)
    args = ap.parse_args()
    salida = args.salida or args.rechazos

    wb = load_workbook(args.rechazos)
    aplicadas = 0
    for nombre in wb.sheetnames:
        if not nombre.lower().startswith("rechazos"):
            continue
        ws = wb[nombre]
        for letra, ancho in ANCHOS_FIJOS.items():
            ws.column_dimensions[letra].width = ancho
        aplicadas += 1
        print(f"  -> anchos fijados en '{nombre}'")

    wb.save(salida)
    print(f"\nTotal hojas formateadas: {aplicadas}")
    print(f"Guardado: {salida}")


if __name__ == "__main__":
    main()
