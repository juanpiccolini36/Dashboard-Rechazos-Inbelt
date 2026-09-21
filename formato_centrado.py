"""
formato_centrado.py
------------------------------------------------------------
Centra horizontal y verticalmente todas las celdas con datos (header
+ filas) en las pestañas "Rechazos *" y las hojas de configuracion,
y activa "Ajustar texto".

El alto de fila NO se fija a mano (salvo la fila 1, que ya quedo en
30 por un pedido anterior) - se deja que Excel lo calcule solo segun
el contenido apenas se abre el archivo, que es como funciona el
ajuste automatico de alto en Excel.

USO:
    python formato_centrado.py --rechazos ARCHIVO.xlsx --salida ARCHIVO_NUEVO.xlsx
"""

import argparse
from openpyxl import load_workbook
from openpyxl.styles import Alignment

HOJAS_A_FORMATEAR_ADEMAS_DE_RECHAZOS = ["Config_Maestro_Pesos", "Config_Tipos_Falla"]
ALTO_FILA_1 = 30


def centrar_hoja(ws):
    alineacion = Alignment(horizontal="center", vertical="center", wrap_text=True)
    headers = [c.value for c in ws[1]]
    n_cols = len(headers)
    for fila in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=n_cols):
        for celda in fila:
            celda.alignment = alineacion

    # la fila 1 mantiene su alto fijo (pedido anterior); el resto se
    # deja en automatico, sin altura fija, para que el ajuste de texto
    # se autocalcule al abrir
    ws.row_dimensions[1].height = ALTO_FILA_1
    for r in range(2, ws.max_row + 1):
        if r in ws.row_dimensions:
            ws.row_dimensions[r].height = None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--rechazos", required=True)
    ap.add_argument("--salida", default=None)
    args = ap.parse_args()
    salida = args.salida or args.rechazos

    wb = load_workbook(args.rechazos)
    aplicadas = 0
    for nombre in wb.sheetnames:
        es_rechazos = nombre.lower().startswith("rechazos")
        es_config = nombre in HOJAS_A_FORMATEAR_ADEMAS_DE_RECHAZOS
        if not (es_rechazos or es_config):
            continue
        ws = wb[nombre]
        centrar_hoja(ws)
        aplicadas += 1
        print(f"  -> celdas centradas + ajuste de texto en '{nombre}'")

    wb.save(salida)
    print(f"\nTotal hojas formateadas: {aplicadas}")
    print(f"Guardado: {salida}")


if __name__ == "__main__":
    main()
